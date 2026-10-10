"""Document ingestion pipeline supporting Markdown, TXT, PDF, and DOCX.

Features:
- Content extraction and cleaning
- SHA256 content deduplication
- Navigation/TOC filtering through semantic chunking
- Context-aware embedding using document title + section + passage
- Transactional persistence to PostgreSQL (knowledge_documents & knowledge_chunks)
"""

from dataclasses import dataclass
import hashlib
import logging
from pathlib import Path
from typing import Any
from uuid import uuid4

from sqlalchemy import delete, select

from src.agent.config import agent_config
from src.agent.rag.chunking import (
    RAG_PIPELINE_VERSION,
    chunk_document,
    embedding_text_for_chunk,
)
from src.agent.rag.embeddings import BaseEmbeddingService, get_embedding_service
from src.db.database import Database
from src.db.models import KnowledgeChunk, KnowledgeDocument, utc_now

logger = logging.getLogger("datt.agent.ingestion")


@dataclass
class IngestionResult:
    """Report of a document ingestion run."""
    document_id: str
    title: str
    source: str
    document_type: str
    chunks_created: int
    status: str  # "created", "updated", "skipped", "failed"
    error: str | None = None


class IngestionService:
    """Handles parsing, chunking, embedding, and storing documents."""

    def __init__(self, db: Database | None = None, embedding_service: BaseEmbeddingService | None = None) -> None:
        self.db = db or agent_config.get_database()
        self._owns_db = db is None
        self.embedder = embedding_service or get_embedding_service()

    def close(self) -> None:
        if self._owns_db:
            self.db.dispose()

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()

    @staticmethod
    def compute_hash(text: str) -> str:
        """Calculate SHA-256 hash of text content."""
        return hashlib.sha256(text.encode("utf-8")).hexdigest()

    def extract_text_from_file(self, file_path: Path | str) -> tuple[str, str]:
        """Extract plain text and determine document type from a file path."""
        path = Path(file_path).resolve()
        if not path.is_file():
            raise FileNotFoundError(f"Document file not found: {path}")

        suffix = path.suffix.lower()
        if suffix in (".md", ".markdown"):
            text = path.read_text(encoding="utf-8", errors="replace")
            return text, "markdown"
        if suffix in (".txt", ".log", ".rst"):
            text = path.read_text(encoding="utf-8", errors="replace")
            return text, "txt"
        if suffix == ".pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(str(path))
                pages: list[str] = []
                for i, page in enumerate(reader.pages, start=1):
                    p_text = page.extract_text() or ""
                    pages.append(f"<!-- page {i} -->\n" + p_text)
                return "\n\n".join(pages), "pdf"
            except Exception as exc:
                raise ValueError(f"Failed to extract PDF text from {path}: {exc}") from exc
        if suffix in (".docx", ".doc"):
            try:
                import docx
                doc = docx.Document(str(path))
                text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
                return text, "docx"
            except Exception as exc:
                raise ValueError(f"Failed to extract DOCX text from {path}: {exc}") from exc
        raise ValueError(f"Unsupported document format: {suffix}")

    def ingest_text(
        self,
        title: str,
        text: str,
        source: str = "manual_entry",
        document_type: str = "manual",
        metadata: dict[str, Any] | None = None,
        force_update: bool = False,
    ) -> IngestionResult:
        """Ingest raw text content into the knowledge base."""
        cleaned = text.strip()
        if not cleaned:
            return IngestionResult(
                document_id="",
                title=title,
                source=source,
                document_type=document_type,
                chunks_created=0,
                status="failed",
                error="Document content is empty",
            )

        content_hash = self.compute_hash(cleaned)
        doc_meta = dict(metadata or {})
        doc_meta.update({
            "source": source,
            "title": title,
            "document_type": document_type,
            "embedding_model": agent_config.embedding_model,
            "embedding_dim": agent_config.embedding_dim,
            "rag_pipeline_version": RAG_PIPELINE_VERSION,
        })

        # 1. Chunk document and drop structural/navigation-only passages.
        chunks = chunk_document(cleaned, doc_metadata=doc_meta)
        if not chunks:
            return IngestionResult(
                document_id="",
                title=title,
                source=source,
                document_type=document_type,
                chunks_created=0,
                status="failed",
                error="No answerable chunks produced from text",
            )

        # 2. Embed a retrieval representation that includes title + active section.
        # Raw chunk.content remains unchanged for citations and grounded answers.
        embedding_inputs = [embedding_text_for_chunk(c, title) for c in chunks]
        embeddings = self.embedder.embed_documents(embedding_inputs)

        for emb in embeddings:
            if emb is not None and len(emb) != agent_config.embedding_dim:
                raise ValueError(
                    f"Incompatible embedding dimension detected: model produced {len(emb)}, "
                    f"but database schema expects {agent_config.embedding_dim}."
                )

        # 3. Short atomic database write transaction.
        with self.db.transaction() as session:
            existing_by_hash = session.scalar(
                select(KnowledgeDocument).where(KnowledgeDocument.content_hash == content_hash)
            )

            existing_meta = (existing_by_hash.doc_metadata or {}) if existing_by_hash else {}
            existing_model = existing_meta.get("embedding_model")
            existing_dim = existing_meta.get("embedding_dim")
            existing_pipeline = existing_meta.get("rag_pipeline_version")
            compatible_existing = (
                existing_by_hash is not None
                and existing_model == agent_config.embedding_model
                and existing_dim == agent_config.embedding_dim
                and existing_pipeline == RAG_PIPELINE_VERSION
            )
            if compatible_existing and not force_update:
                logger.info(
                    "Document '%s' identical hash already exists with current embedding/RAG pipeline; skipping",
                    title,
                )
                return IngestionResult(
                    document_id=str(existing_by_hash.id),
                    title=existing_by_hash.title,
                    source=existing_by_hash.source,
                    document_type=existing_by_hash.document_type,
                    chunks_created=existing_by_hash.chunk_count,
                    status="skipped",
                )

            existing_by_source = session.scalar(
                select(KnowledgeDocument).where(KnowledgeDocument.source == source)
            )

            if existing_by_source:
                doc_record = existing_by_source
                doc_record.title = title
                doc_record.document_type = document_type
                doc_record.content_hash = content_hash
                doc_record.chunk_count = len(chunks)
                doc_record.doc_metadata = doc_meta
                doc_record.updated_at = utc_now()
                session.execute(
                    delete(KnowledgeChunk).where(KnowledgeChunk.document_id == doc_record.id)
                )
                status = "updated"
            else:
                doc_record = KnowledgeDocument(
                    id=uuid4(),
                    title=title,
                    source=source,
                    document_type=document_type,
                    content_hash=content_hash,
                    chunk_count=len(chunks),
                    doc_metadata=doc_meta,
                    created_at=utc_now(),
                    updated_at=utc_now(),
                )
                session.add(doc_record)
                status = "created"

            session.flush()

            for i, chunk in enumerate(chunks):
                emb = embeddings[i] if i < len(embeddings) else None
                chunk_meta = dict(chunk.metadata)
                chunk_meta.update({
                    "embedding_model": agent_config.embedding_model,
                    "embedding_dim": agent_config.embedding_dim,
                    "rag_pipeline_version": RAG_PIPELINE_VERSION,
                })
                chunk_rec = KnowledgeChunk(
                    id=uuid4(),
                    document_id=doc_record.id,
                    chunk_index=chunk.chunk_index,
                    content=chunk.content,
                    embedding=emb,
                    chunk_metadata=chunk_meta,
                    created_at=utc_now(),
                    updated_at=utc_now(),
                )
                session.add(chunk_rec)

            session.flush()

            logger.info("Ingested '%s': %d chunks (%s)", title, len(chunks), status)
            return IngestionResult(
                document_id=str(doc_record.id),
                title=doc_record.title,
                source=doc_record.source,
                document_type=doc_record.document_type,
                chunks_created=len(chunks),
                status=status,
            )

    def ingest_file(self, file_path: Path | str, title: str | None = None, force_update: bool = False) -> IngestionResult:
        """Extract text from file and ingest into knowledge base."""
        path = Path(file_path).resolve()
        doc_title = title or path.stem.replace("_", " ").title()
        text, doc_type = self.extract_text_from_file(path)
        return self.ingest_text(
            title=doc_title,
            text=text,
            source=str(path),
            document_type=doc_type,
            metadata={"filename": path.name, "file_size": path.stat().st_size},
            force_update=force_update,
        )

    def ingest_directory(
        self,
        directory_path: Path | str,
        glob_patterns: list[str] | None = None,
        force_update: bool = False,
    ) -> list[IngestionResult]:
        """Ingest all supported documents in a directory."""
        dir_p = Path(directory_path).resolve()
        patterns = glob_patterns or ["*.md", "*.txt", "*.pdf", "*.docx"]
        results: list[IngestionResult] = []

        files: list[Path] = []
        for pattern in patterns:
            files.extend(dir_p.glob(pattern))

        for file_p in sorted(set(files)):
            try:
                res = self.ingest_file(file_p, force_update=force_update)
                results.append(res)
            except Exception as exc:
                logger.error("Failed to ingest file %s: %s", file_p, exc)
                results.append(
                    IngestionResult(
                        document_id="",
                        title=file_p.stem,
                        source=str(file_p),
                        document_type=file_p.suffix.lstrip("."),
                        chunks_created=0,
                        status="failed",
                        error=str(exc),
                    )
                )
        return results
